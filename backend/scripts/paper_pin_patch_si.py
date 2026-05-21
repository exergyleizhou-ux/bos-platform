# -*- coding: utf-8 -*-
"""SI patches — bring BOS_Paper1_JCP_SI.docx to submission-ready state.

Conservative same-pattern as scratch_y1_patch.py:
- paragraph.text find-and-replace only
- hit-count guard per patch
- backup before any write
- save in-place (SI has no gate test; safe to overwrite)
"""
from pathlib import Path
import hashlib
import shutil

from docx import Document

SI = Path(
    r"C:\Users\10420\Desktop\bos 0506\Paper1_BT"
    r"\BOS_Paper1_JCP_SI.docx"
)
BACKUP = SI.parent / "BOS_Paper1_JCP_SI.docx.pre-si-patch"


def count_hits(doc, find_str):
    return sum(1 for p in doc.paragraphs if find_str in p.text)


def sha256_of_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def apply(doc, find_str, replace_str, name, expected_hits):
    hits = count_hits(doc, find_str)
    if hits != expected_hits:
        print(f"  [{name}] SKIP: expected {expected_hits} hit(s), "
              f"found {hits} for {find_str!r}")
        return False
    replaced_in = 0
    for p in doc.paragraphs:
        if find_str in p.text:
            p.text = p.text.replace(find_str, replace_str)
            replaced_in += 1
    print(f"  [{name}] OK: replaced in {replaced_in} paragraph(s)")
    return True


def main():
    print(f"Loading: {SI}")
    print(f"Original SHA: {sha256_of_file(SI)}")
    print(f"Original size: {SI.stat().st_size} bytes")

    # Backup
    shutil.copy(SI, BACKUP)
    print(f"Backup: {BACKUP}")
    print(f"Backup SHA: {sha256_of_file(BACKUP)}")
    print()

    doc = Document(str(SI))
    print(f"Paragraphs: {len(doc.paragraphs)}")

    # Recon all "placeholder" / "to be issued" / OSF / Zenodo hits with
    # context, so we know which to patch.
    print()
    print("=" * 70)
    print("SI Recon — finding patchable placeholders")
    print("=" * 70)
    placeholder_keywords = [
        "[OSF DOI to be issued]",
        "OSF [DOI to be issued]",
        "[Zenodo DOI to be issued]",
        "Zenodo (DOI: [Zenodo DOI to be issued])",
        "[URL to be issued]",
        "[hash to be issued]",
        "commit [hash to be issued]",
        "Zenodo [DOI placeholder]",
        "[DOI placeholder]",
        "[commit: XXXXXXX]",
        "[tag: bt-submission-v1]",
    ]
    for kw in placeholder_keywords:
        h = count_hits(doc, kw)
        marker = ""
        if h == 1:
            marker = " <-- UNIQUE"
        elif h > 1:
            marker = f" <-- {h}x"
        elif h == 0:
            marker = " (absent)"
        print(f"  {kw!r:50s}: {h} hit(s){marker}")
        if 1 <= h <= 3:
            for i, p in enumerate(doc.paragraphs):
                if kw in p.text:
                    print(f"    para[{i}]: {p.text[:160]!r}")

    print()
    print("=" * 70)
    print("SI Patch — applying replacements")
    print("=" * 70)

    # Patch SI.1 — Zenodo DOI placeholder (Software Note §S_Software-5)
    apply(doc,
          find_str="[Zenodo DOI to be issued]",
          replace_str=("[Zenodo DOI: 10.5281/zenodo.XXXXXXX — pending "
                       "mint after GitHub Release publication]"),
          name="SI.1 Zenodo DOI placeholder",
          expected_hits=1)

    # Patch SI.2 — URL placeholder
    apply(doc,
          find_str="[URL to be issued]",
          replace_str="https://github.com/exergyleizhou-ux/bos-platform",
          name="SI.2 GitHub URL placeholder",
          expected_hits=1)

    # Patch SI.3 — commit hash placeholder (BT-style)
    apply(doc,
          find_str="[commit: XXXXXXX]",
          replace_str="commit: 92a7a0f",
          name="SI.3 commit hash placeholder",
          expected_hits=1)

    # Patch SI.4 — tag placeholder (BT-style legacy)
    apply(doc,
          find_str="[tag: bt-submission-v1]",
          replace_str="tag: v0.9.0-paper1",
          name="SI.4 tag placeholder",
          expected_hits=1)

    # Patch SI.5 — OSF DOI single-bracket form
    apply(doc,
          find_str="[OSF DOI to be issued]",
          replace_str=("[OSF DOI: 10.17605/OSF.IO/XXXXX — pending "
                       "pre-registration submission]"),
          name="SI.5 OSF DOI single",
          expected_hits=1)

    # Patch SI.6 — Zenodo Data Avail anchor (if same as main paper)
    apply(doc,
          find_str="deposited at Zenodo [DOI placeholder].",
          replace_str=("deposited at Zenodo [DOI: 10.5281/zenodo.XXXXXXX "
                       "— pending mint]."),
          name="SI.6 Zenodo Data Avail",
          expected_hits=1)

    # Patch SI.7 — Zenodo [DOI placeholder] anchor (SI-specific form)
    apply(doc,
          find_str="Zenodo [DOI placeholder]",
          replace_str=("Zenodo [DOI: 10.5281/zenodo.XXXXXXX — pending "
                       "mint]"),
          name="SI.7 Zenodo bracketed placeholder",
          expected_hits=1)

    # Save in-place
    doc.save(str(SI))

    new_sha = sha256_of_file(SI)
    new_size = SI.stat().st_size
    print()
    print("=" * 70)
    print("Saved SI")
    print("=" * 70)
    print(f"  Path: {SI}")
    print(f"  New SHA: {new_sha}")
    print(f"  New size: {new_size} bytes")
    print(f"  Backup at: {BACKUP}")

    # Post-patch sanity
    patched = Document(str(SI))
    print()
    print("Post-patch sanity (placeholders should be 0 hits):")
    for find, name in [
        ("[Zenodo DOI to be issued]", "SI.1"),
        ("[URL to be issued]", "SI.2"),
        ("[commit: XXXXXXX]", "SI.3"),
        ("[tag: bt-submission-v1]", "SI.4"),
        ("[OSF DOI to be issued]", "SI.5"),
        ("deposited at Zenodo [DOI placeholder].", "SI.6"),
    ]:
        h = count_hits(patched, find)
        status = "OK (absent)" if h == 0 else f"FAIL ({h} hits)"
        print(f"  {name}: {status}")


if __name__ == "__main__":
    main()
